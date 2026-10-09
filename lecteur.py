from pathlib import Path
import os
import re
import zipfile
import xml.etree.ElementTree as ET


def _read_docx_fallback(path: Path) -> str:
    """Essaie de lire un DOCX directement si python-docx refuse le fichier."""
    try:
        with zipfile.ZipFile(path, "r") as archive:
            names = set(archive.namelist())
            document_name = "word/document.xml"

            if document_name not in names:
                raise RuntimeError(
                    "Le fichier .docx est incomplet : word/document.xml est absent."
                )

            xml_data = archive.read(document_name)
    except zipfile.BadZipFile as exc:
        raise RuntimeError(
            "Le fichier porte l'extension .docx mais ce n'est pas une archive DOCX valide. "
            "Ouvre-le dans Word puis fais Fichier > Enregistrer sous > Document Word (*.docx)."
        ) from exc
    except KeyError as exc:
        raise RuntimeError("Impossible de trouver le contenu principal du document DOCX.") from exc

    try:
        root = ET.fromstring(xml_data)
    except ET.ParseError as exc:
        raise RuntimeError("Le document DOCX contient un XML illisible.") from exc

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraphs = []

    for paragraph in root.findall(".//w:p", ns):
        parts = [node.text or "" for node in paragraph.findall(".//w:t", ns)]
        text = "".join(parts).strip()
        if text:
            paragraphs.append(text)

    return "\n".join(paragraphs)


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise RuntimeError(
            "Le module python-docx est absent. Vérifie requirements.txt."
        ) from exc

    try:
        doc = Document(str(path))
        chunks = []

        # Garde l'ordre entre paragraphes et tableaux : les consignes et les
        # données d'un exercice peuvent être réparties entre les deux.
        for block in doc.iter_inner_content():
            if isinstance(block, Paragraph):
                if block.text.strip():
                    chunks.append(block.text.strip())
            elif isinstance(block, Table):
                for row in block.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    line = " | ".join(cell for cell in cells if cell)
                    if line:
                        chunks.append(line)

        return "\n".join(chunks)

    except Exception as exc:
        # Certains fichiers .docx sont des archives ZIP lisibles mais dont
        # la structure OOXML est imparfaite. Le fallback permet de récupérer
        # directement word/document.xml.
        try:
            fallback = _read_docx_fallback(path)
            if fallback.strip():
                return fallback
        except Exception as fallback_exc:
            raise RuntimeError(
                "Impossible de lire ce DOCX. Le fichier semble mal formé ou n'est pas un vrai DOCX. "
                "Essaie de l'ouvrir dans Word puis de faire 'Enregistrer sous > .docx'. "
                f"Détail : {fallback_exc}"
            ) from exc

        raise RuntimeError(
            "Le DOCX ne contient aucun texte exploitable. "
            "Essaie de l'ouvrir puis de l'enregistrer à nouveau au format .docx."
        ) from exc




def _read_odt(path: Path) -> str:
    """Lit un fichier OpenDocument Text (.odt) sans dépendance externe.

    Un ODT est une archive ZIP qui contient notamment content.xml.
    Cette méthode récupère les paragraphes, titres et éléments de liste
    dans leur ordre de lecture.
    """
    try:
        with zipfile.ZipFile(path, "r") as archive:
            names = set(archive.namelist())
            if "content.xml" not in names:
                raise RuntimeError(
                    "Le fichier .odt est incomplet : content.xml est absent."
                )
            xml_data = archive.read("content.xml")
    except zipfile.BadZipFile as exc:
        raise RuntimeError(
            "Le fichier porte l'extension .odt mais n'est pas une archive ODT valide. "
            "Ouvre-le dans LibreOffice/Word puis fais Enregistrer sous > .odt."
        ) from exc
    except KeyError as exc:
        raise RuntimeError("Impossible de trouver le contenu principal du document ODT.") from exc

    try:
        root = ET.fromstring(xml_data)
    except ET.ParseError as exc:
        raise RuntimeError("Le document ODT contient un XML illisible.") from exc

    ns = {
        "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
        "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
        "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    }

    body = root.find(".//office:body/office:text", ns)
    if body is None:
        raise RuntimeError("Le document ODT ne contient pas de texte exploitable.")

    chunks = []

    def element_text(element):
        parts = []
        if element.text:
            parts.append(element.text)
        for child in element:
            parts.append(element_text(child))
            if child.tail:
                parts.append(child.tail)
        return "".join(parts)

    def collect(element):
        tag = element.tag
        if tag in {f"{{{ns['text']}}}h", f"{{{ns['text']}}}p"}:
            value = element_text(element).strip()
            if value:
                chunks.append(value)
            return
        elif tag == f"{{{ns['table']}}}table":
            for row in element.findall(".//table:table-row", ns):
                cells = []
                for cell in row.findall("table:table-cell", ns):
                    cell_text = " ".join(
                        value.strip()
                        for value in (
                            element_text(paragraph)
                            for paragraph in cell.findall(".//text:p", ns)
                        )
                        if value.strip()
                    )
                    if cell_text:
                        cells.append(cell_text)
                if cells:
                    chunks.append(" | ".join(cells))
            return
        for child in element:
            collect(child)

    for child in body:
        collect(child)

    return "\n".join(chunks).strip()

def read_file(path: Path, *, exercise_number=None) -> str:
    ext = path.suffix.lower()

    if ext == ".txt":
        content = path.read_bytes()
        for encoding in ("utf-8-sig", "utf-16", "cp1252"):
            try:
                return content.decode(encoding).strip()
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace").strip()

    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            text = "\n\n".join(
                (page.extract_text() or "")
                for page in reader.pages
            )
            return text
        except Exception as exc:
            raise RuntimeError(f"Impossible de lire le PDF : {exc}") from exc

    if ext == ".docx":
        return _read_docx(path)

    if ext == ".odt":
        return _read_odt(path)

    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(path))
        chunks = []
        for slide in prs.slides:
            for shape in slide.shapes:
                text = getattr(shape, "text", "")
                if text:
                    chunks.append(text)
        return "\n".join(chunks)

    if ext in {".jpg", ".jpeg", ".png", ".webp"}:
        try:
            import pytesseract
            from PIL import Image, ImageEnhance, ImageFilter, ImageOps
        except ImportError as exc:
            raise RuntimeError(
                "OCR indisponible : installe Pillow et pytesseract."
            ) from exc

        try:
            image = ImageOps.exif_transpose(Image.open(path))
            image.load()

            max_side = 3000
            largest_side = max(image.size)
            if largest_side > max_side:
                ratio = max_side / largest_side
                image = image.resize(
                    (
                        max(1, int(image.width * ratio)),
                        max(1, int(image.height * ratio)),
                    ),
                    Image.Resampling.LANCZOS,
                )

            if image.mode not in {"RGB", "L"}:
                image = image.convert("RGB")

            gray = ImageOps.grayscale(image)
            gray = ImageOps.autocontrast(gray)
            gray = ImageEnhance.Sharpness(gray).enhance(1.7)
            gray = gray.filter(ImageFilter.MedianFilter(size=3))

        except Exception as exc:
            raise RuntimeError(
                f"Impossible d'ouvrir l'image : {exc}"
            ) from exc

        tesseract_cmd = os.getenv("TESSERACT_CMD", "").strip()
        if not tesseract_cmd and os.name == "nt":
            install_roots = [
                Path(os.getenv("ProgramFiles", r"C:\Program Files")),
                Path(os.getenv("LOCALAPPDATA", str(Path.home() / "AppData" / "Local"))),
            ]
            for root in install_roots:
                candidate = root / "Tesseract-OCR" / "tesseract.exe"
                if candidate.is_file():
                    tesseract_cmd = str(candidate)
                    break
        if tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

        languages = []
        try:
            installed = set(pytesseract.get_languages(config=""))
            for language in ("fra", "ita", "eng"):
                if language in installed:
                    languages.append(language)
        except Exception:
            languages = ["fra", "ita", "eng"]

        lang = "+".join(languages) if languages else "eng"

        try:
            selected_image = gray
            if exercise_number is not None:
                data = pytesseract.image_to_data(
                    gray,
                    lang=lang,
                    config="--oem 3 --psm 11",
                    output_type=pytesseract.Output.DICT,
                )
                headings = []
                min_width = max(20, int(gray.width * 0.02))
                min_height = max(16, int(gray.height * 0.01))
                for index, raw_word in enumerate(data.get("text", [])):
                    digits = re.sub(r"\D", "", str(raw_word))
                    if not digits or len(digits) > 3:
                        continue
                    left = int(data["left"][index])
                    top = int(data["top"][index])
                    width = int(data["width"][index])
                    height = int(data["height"][index])
                    try:
                        confidence = float(data["conf"][index])
                    except (TypeError, ValueError):
                        confidence = -1
                    if (
                        left < gray.width * 0.55
                        and width >= min_width
                        and height >= min_height
                        and confidence >= 15
                    ):
                        headings.append((int(digits), left, top, width, height, confidence))

                requested_number = int(exercise_number)
                requested_heading = next(
                    (item for item in headings if item[0] == requested_number),
                    None,
                )
                if headings and requested_heading is None:
                    raise RuntimeError(
                        f"Je ne repère pas le numéro {requested_number} sur cette photo. "
                        "Pour éviter de résoudre un exercice voisin, encadre seulement "
                        "l’exercice demandé et renvoie la photo."
                    )

                if requested_heading:
                    _, left, top, _, height, _ = requested_heading
                    next_heading = next(
                        (
                            item for item in sorted(headings, key=lambda value: value[2])
                            if item[2] > top + height
                            and item[0] > requested_number
                        ),
                        None,
                    )
                    padding_x = max(12, int(gray.width * 0.025))
                    padding_y = max(10, int(gray.height * 0.008))
                    crop_top = max(0, top - padding_y)
                    crop_bottom = (
                        max(crop_top + 1, next_heading[2] - padding_y)
                        if next_heading
                        else gray.height
                    )
                    crop_left = max(0, left - padding_x)
                    selected_image = gray.crop((crop_left, crop_top, gray.width, crop_bottom))

            # Deux modes OCR : texte en blocs et mise en page plus libre.
            results = []
            for psm in (6, 11):
                text = pytesseract.image_to_string(
                    selected_image,
                    lang=lang,
                    config=f"--oem 3 --psm {psm}",
                ).strip()
                if text:
                    results.append(text)
        except pytesseract.TesseractNotFoundError as exc:
            raise RuntimeError(
                "Tesseract OCR est introuvable. Installe Tesseract OCR, puis indique "
                "son chemin dans TESSERACT_CMD (par exemple : "
                "C:\\Program Files\\Tesseract-OCR\\tesseract.exe) et redémarre RéviAI."
            ) from exc
        except pytesseract.TesseractError as exc:
            raise RuntimeError(
                "Tesseract n'a pas pu lire cette image. Vérifie que les données de langue "
                "française (fra) ou anglaise (eng) sont installées. "
                f"Détail : {exc}"
            ) from exc

        if not results:
            raise RuntimeError(
                "Tesseract n'a détecté aucun texte. Essaie une photo plus nette, bien cadrée et éclairée."
            )

        # Prend généralement le résultat le plus complet.
        return max(results, key=len)

    raise RuntimeError(f"Format non supporté : {ext or 'inconnu'}")
